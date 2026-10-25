"""Shared upload/API normalization. No model or uploaded content can execute code."""

import csv
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import date, datetime

from fastapi import HTTPException

from decisioniq.dashboard import store

MAX_ROWS = 25_000
CURRENCIES = {"GBP", "USD", "EUR", "INR", "AUD", "CAD", "JPY"}
ALIASES = {
    "record_id": ["recordid", "transactionid", "eventid", "id"],
    "date": ["date", "invoicedate", "transactiondate", "timestamp", "occurredat", "ds"],
    "sales": ["sales", "salesvalue", "revenue", "amount", "total", "y"],
    "product_id": ["productid", "stockcode", "sku", "itemid"],
    "product": ["product", "description", "productname", "item"],
    "units": ["units", "quantity", "qty"],
    "unit_price": ["unitprice", "price"],
    "currency": ["currency", "currencycode"],
}


def suggestions(columns):
    normalized = {re.sub(r"[^a-z0-9]", "", column.lower()): column for column in columns}
    return {
        key: next((normalized[value] for value in values if value in normalized), "")
        for key, values in ALIASES.items()
    }


def clean_cell(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return str(value).strip() if not isinstance(value, (int, float, bool)) else value


def read_upload(content, filename):
    if not content or len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Upload a nonempty file no larger than 5 MB.")
    suffix = filename.rsplit(".", 1)[-1].lower()
    try:
        if suffix == "csv":
            text = content.decode("utf-8-sig")
            dialect = (
                csv.Sniffer().sniff(text[:8192], delimiters=",;\t") if "\n" in text else csv.excel
            )
            reader = csv.DictReader(io.StringIO(text), dialect=dialect)
            columns = reader.fieldnames or []
            records = []
            for index, row in enumerate(reader):
                if index >= MAX_ROWS:
                    raise HTTPException(413, f"Use at most {MAX_ROWS:,} rows per batch.")
                if None in row:
                    raise ValueError("Some rows have more cells than the header.")
                records.append({k: clean_cell(v) for k, v in row.items()})
        elif suffix == "xlsx":
            from openpyxl import load_workbook

            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("The expanded workbook is too large; use a smaller export.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                if sheet is None or sheet.max_column > 100 or sheet.max_row > MAX_ROWS + 1:
                    raise ValueError("Use at most 100 columns and 25,000 rows.")
                iterator = sheet.iter_rows(values_only=True)
                columns = [
                    str(value).strip() if value is not None else "" for value in next(iterator)
                ]
                records = [
                    dict(zip(columns, map(clean_cell, row)))
                    for row in iterator
                    if any(v is not None for v in row)
                ]
            finally:
                workbook.close()
        else:
            raise HTTPException(
                415, "Supported uploads are .csv and .xlsx. The integration API accepts JSON."
            )
        if (
            not columns
            or any(not name or len(name) > 100 for name in columns)
            or len(set(columns)) != len(columns)
        ):
            raise ValueError("Use unique, nonempty column headers of at most 100 characters.")
        if len(columns) > 100 or not records:
            raise ValueError("The file must contain data and no more than 100 columns.")
        if any(isinstance(v, str) and len(v) > 5000 for row in records for v in row.values()):
            raise ValueError("A cell exceeds the 5,000-character limit.")
        return columns, records
    except HTTPException:
        raise
    except (
        ValueError,
        UnicodeDecodeError,
        csv.Error,
        zipfile.BadZipFile,
        KeyError,
        StopIteration,
    ) as exc:
        raise HTTPException(422, f"Could not read this file: {exc}") from exc


def create_batch(company_id, name, records, source="upload", options=None):
    batch_id = store.uid()
    columns = list(dict.fromkeys(key for row in records[:100] for key in row))
    config = options or {"mapping": suggestions(columns)}
    with store.connect() as connection:
        connection.execute(
            "INSERT INTO batches VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (
                batch_id,
                company_id,
                name[:160],
                source,
                "sales",
                "needs_mapping",
                store.stamp(),
                None,
                json.dumps(records),
                json.dumps(config),
                "{}",
            ),
        )
        store.activity(
            company_id,
            "upload",
            "Sales data received",
            f"{name[:100]} · {len(records):,} rows · review required",
            connection,
        )
    return batch_id


def get_batch(company_id, batch_id, include_raw=False):
    result = store.rows("SELECT * FROM batches WHERE company_id=? AND id=?", (company_id, batch_id))
    if not result:
        raise HTTPException(404, "Batch not found in this workspace.")
    batch = result[0]
    records = json.loads(batch.pop("raw"))
    batch["options"] = json.loads(batch["options"])
    batch["result"] = json.loads(batch["result"])
    batch["row_count"] = len(records)
    batch["columns"] = list(dict.fromkeys(key for row in records[:100] for key in row))
    batch["preview"] = records[:5]
    if include_raw:
        batch["records"] = records
    return batch


def numeric(value, field, number_format="decimal_point"):
    if value is None or value == "":
        raise ValueError(f"{field} is missing")
    if isinstance(value, bool):
        raise ValueError(f"{field} must be numeric")
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        text = str(value).strip()
        pattern = (
            r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?:[eE][+-]?\d+)?"
            if number_format == "decimal_point"
            else r"[+-]?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?"
        )
        if not re.fullmatch(pattern, text):
            raise ValueError(
                f"{field} has an ambiguous number format; select the correct decimal convention and remove currency symbols"
            )
        result = (
            float(text.replace(",", ""))
            if number_format == "decimal_point"
            else float(text.replace(".", "").replace(",", "."))
        )
    if not math.isfinite(result) or abs(result) > 1e12:
        raise ValueError(f"{field} must be finite and within ±1 trillion")
    return result


def normalize(company, records, options):
    import pandas as pd

    mapping = options.get("mapping", {})
    currency = options.get("currency", company["currency"]).upper()
    if currency not in CURRENCIES or currency != company["currency"]:
        raise HTTPException(
            422,
            f"Use this workspace's {company['currency']} currency. Create a separate workspace for another currency; no automatic FX conversion is applied.",
        )
    if not mapping.get("date") or not (
        mapping.get("sales") or (mapping.get("units") and mapping.get("unit_price"))
    ):
        return (
            [],
            [
                {
                    "row": 0,
                    "message": "Map a date and either a sales amount or both units and unit price.",
                }
            ],
            [],
        )
    normalized, errors = [], []
    warnings = []
    if not mapping.get("record_id"):
        warnings.append(
            "No record ID is mapped. Exact normalized rows will be deduplicated. Map stable transaction IDs when identical purchases can occur."
        )
    for index, raw in enumerate(records, 2):

        def value(key):
            return raw.get(mapping.get(key, ""))

        try:
            raw_date = value("date")
            if raw_date in {None, ""}:
                raise ValueError("date is missing")
            mode = options.get("date_format", "iso")
            if mode == "iso" and not re.match(r"^\d{4}-\d{2}-\d{2}(?:$|[T ])", str(raw_date)):
                raise ValueError("Select the file's date order; ISO dates must use YYYY-MM-DD")
            parsed = pd.to_datetime(str(raw_date), dayfirst=mode == "day_first", errors="raise")
            if parsed.tzinfo:
                parsed = parsed.tz_convert(company["timezone"])
            record_date = parsed.date()
            from zoneinfo import ZoneInfo

            if record_date > store.now().astimezone(ZoneInfo(company["timezone"])).date():
                raise ValueError("Observed sales cannot have a future transaction date")
            units = (
                numeric(value("units"), "units", options.get("number_format", "decimal_point"))
                if mapping.get("units") and value("units") not in {None, ""}
                else None
            )
            if not mapping.get("sales") and units is None:
                raise ValueError("units is missing; provide it to calculate sales from unit price")
            sales = (
                numeric(value("sales"), "sales", options.get("number_format", "decimal_point"))
                if mapping.get("sales")
                else units
                * numeric(
                    value("unit_price"), "unit price", options.get("number_format", "decimal_point")
                )
            )
            if value("currency") and str(value("currency")).upper() != currency:
                raise ValueError(
                    f"currency does not match {currency}; separate mixed-currency data before processing"
                )
            product_name = str(value("product") or value("product_id") or "All products").strip()[
                :240
            ]
            product_id = value("product_id")
            if not product_id:
                product_id = (
                    "name:"
                    + hashlib.sha256(
                        " ".join(product_name.casefold().split()).encode()
                    ).hexdigest()[:16]
                    if value("product")
                    else "all"
                )
            item = {
                "date": record_date.isoformat(),
                "product_id": str(product_id)[:100],
                "product": product_name,
                "units": units,
                "sales": round(sales, 6),
                "currency": currency,
            }
            record_id = value("record_id")
            item["record_id"] = (
                (
                    str(record_id)
                    if str(record_id).startswith(("external:", "content:"))
                    else "external:" + str(record_id)
                )
                if record_id not in {None, ""}
                else "content:"
                + hashlib.sha256(json.dumps(item, sort_keys=True).encode()).hexdigest()
            )
            normalized.append(item)
        except (ValueError, TypeError, OverflowError) as exc:
            errors.append({"row": index, "message": str(exc)[:240]})
    return normalized, errors, warnings


def validate_batch(company, batch_id, options):
    batch = get_batch(company["id"], batch_id, include_raw=True)
    normalized, errors, warnings = normalize(company, batch["records"], options)
    # An existing stable identifier may only be replayed with the same values.
    existing = {
        row["record_id"]: row
        for row in store.rows("SELECT * FROM sales WHERE company_id=?", (company["id"],))
    }
    seen = {}
    for index, item in enumerate(normalized, 2):
        other = existing.get(item["record_id"]) or seen.get(item["record_id"])
        if other and any(
            other.get(key) != item[key]
            for key in ["date", "product_id", "product", "units", "sales", "currency"]
        ):
            errors.append(
                {
                    "row": index,
                    "message": "This record ID already exists with different values. Correct the source or provide a distinct ID; existing observations are not overwritten.",
                }
            )
        seen[item["record_id"]] = item
    response = {
        "valid_rows": len(normalized),
        "error_count": len(errors),
        "errors": errors[:20],
        "warnings": warnings,
        "preview": normalized[:5],
        "date_start": min((row["date"] for row in normalized), default=None),
        "date_end": max((row["date"] for row in normalized), default=None),
    }
    return normalized, response


def process_batch(company_id, batch_id):
    import logging
    from decisioniq.dashboard.analytics import fingerprint, refresh_forecast
    from decisioniq.dashboard.reminders import evaluate_reminders

    company = store.company(company_id)
    batch = get_batch(company_id, batch_id)
    try:
        normalized, validation = validate_batch(company, batch_id, batch["options"])
        if validation["error_count"]:
            raise ValueError(validation["errors"][0]["message"])
        inserted = 0
        with store.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            for item in normalized:
                result = connection.execute(
                    "INSERT OR IGNORE INTO sales VALUES (?,?,?,?,?,?,?,?,?)",
                    (
                        company_id,
                        item["record_id"],
                        item["date"],
                        item["product_id"],
                        item["product"],
                        item["units"],
                        item["sales"],
                        item["currency"],
                        batch_id,
                    ),
                )
                if not result.rowcount:
                    other = dict(
                        connection.execute(
                            "SELECT * FROM sales WHERE company_id=? AND record_id=?",
                            (company_id, item["record_id"]),
                        ).fetchone()
                    )
                    if any(
                        other[key] != item[key]
                        for key in ["date", "product_id", "product", "units", "sales", "currency"]
                    ):
                        raise ValueError(
                            "A concurrent batch used this record ID with different values. Existing observations were preserved."
                        )
                inserted += result.rowcount
            summary = {
                "inserted": inserted,
                "duplicates": len(normalized) - inserted,
                "warnings": validation["warnings"],
            }
            connection.execute(
                "UPDATE batches SET result=? WHERE id=? AND company_id=?",
                (json.dumps(summary), batch_id, company_id),
            )
        store.update_company(
            company_id, calendar_complete=batch["options"].get("calendar_complete", False)
        )
    except Exception as exc:
        logging.getLogger(__name__).exception("Batch processing failed")
        with store.connect() as connection:
            connection.execute(
                "UPDATE batches SET status='failed',result=? WHERE id=? AND company_id=?",
                (json.dumps({"error": str(exc)[:300]}), batch_id, company_id),
            )
        return
    # Observation acceptance and forecast availability are different outcomes.
    try:
        forecast = refresh_forecast(company_id)
        summary["forecast_status"] = forecast["status"]
    except Exception:
        logging.getLogger(__name__).exception("Forecast preparation failed")
        summary["forecast_status"] = "unavailable"
        unavailable = {
            "status": "unavailable",
            "fingerprint": fingerprint(company_id),
            "trained_at": store.stamp(),
            "future": [],
            "model": None,
            "message": "Sales records were saved, but forecast preparation failed. The next processed batch will retry analysis.",
        }
        with store.connect() as connection:
            connection.execute(
                "INSERT INTO forecasts VALUES (?,?) ON CONFLICT(company_id) DO UPDATE SET value=excluded.value",
                (company_id, json.dumps(unavailable)),
            )
    with store.connect() as connection:
        connection.execute(
            "UPDATE batches SET status='processed',processed_at=?,result=? WHERE id=? AND company_id=?",
            (store.stamp(), json.dumps(summary), batch_id, company_id),
        )
        store.activity(
            company_id,
            "processed",
            "Sales data processed",
            f"{inserted:,} new observations · {summary['duplicates']:,} duplicates skipped",
            connection,
        )
    evaluate_reminders()
